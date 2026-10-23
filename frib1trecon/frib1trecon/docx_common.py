"""
Estilo visual compartido para los DOCX de frib1trecon. cyber_health_report.py
(resumen ejecutivo de 6 bloques) ya tenia su propio sistema de banners
navy/cajas de color; este modulo porta las mismas piezas (misma paleta
hex, mismo criterio de heading()/shade()) para que informe.docx/
informe_master.docx (el detalle tecnico completo, ver report.py) se vean
como el mismo sistema visual en vez de con el estilo generico de Word por
defecto, sin recortar ni un dato tecnico del contenido existente.
"""
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

NAVY = RGBColor(0x0B, 0x1B, 0x4D)
NAVY_HEX = "0B1B4D"
INK = RGBColor(0x16, 0x21, 0x3F)
GREEN = RGBColor(0x1F, 0x8F, 0x62)
RED = RGBColor(0xC9, 0x34, 0x3A)
MUTED = RGBColor(0x6B, 0x77, 0x94)
WHITE = RGBColor(0xFF, 0xFF, 0xFF)
LIGHT_HEX = "F5F7FC"
HEADING_FONT = "Poppins"  # Word sustituye por una fuente similar si no esta instalada


def shade(cell, hex_color: str) -> None:
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    tcPr.append(shd)


def shade_paragraph(paragraph, hex_color: str) -> None:
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_color)
    pPr.append(shd)


def setup_document(doc, margin_cm: float = 2.2) -> None:
    """Margenes + fuente base, igual que cyber_health_report.py, para que
    ambos DOCX compartan la misma caja de texto y tipografia de cuerpo."""
    for s in doc.sections:
        s.left_margin = s.right_margin = Cm(margin_cm)
        s.top_margin = s.bottom_margin = Cm(margin_cm)
    base = doc.styles["Normal"]
    base.font.name = "Calibri"
    base.font.size = Pt(10)


def heading(doc, text: str, level: int = 1):
    """Nivel 1: banda navy de ancho completo, texto blanco negrita (igual
    que cyber_health_report.py). Nivel 2/3: texto negrita navy/ink, sin
    banda, para no saturar la pagina de bloques de color."""
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(12 if level == 1 else 8)
    p.paragraph_format.space_after = Pt(6 if level == 1 else 4)
    r = p.add_run(text)
    r.bold = True
    r.font.name = HEADING_FONT
    if level == 1:
        shade_paragraph(p, NAVY_HEX)
        pPr = p._p.get_or_add_pPr()
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), "120")
        ind.set(qn("w:right"), "120")
        pPr.append(ind)
        spacing = OxmlElement("w:spacing")
        spacing.set(qn("w:before"), "40")
        spacing.set(qn("w:after"), "40")
        spacing.set(qn("w:line"), "240")
        spacing.set(qn("w:lineRule"), "auto")
        pPr.append(spacing)
        r.font.size = Pt(12.5)
        r.font.color.rgb = WHITE
    elif level == 2:
        r.font.size = Pt(11.5)
        r.font.color.rgb = NAVY
    else:
        r.font.size = Pt(10.5)
        r.font.color.rgb = INK
    return p


def style_table(table, light_rows: bool = True) -> None:
    """Re-estiliza en sitio una tabla ya creada con `doc.add_table` cuya
    fila 0 ya tiene el texto de cabecera: la sombrea navy con texto blanco
    negrita y alterna el fondo de las filas de datos, sin tocar el
    contenido. `table.style` debe ser 'Table Grid' (o similar con bordes
    finos) para que las lineas de cuadricula no choquen con el navy."""
    for c in table.rows[0].cells:
        shade(c, NAVY_HEX)
        for p in c.paragraphs:
            p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            for r in p.runs:
                r.font.color.rgb = WHITE
                r.bold = True
                r.font.size = Pt(9)
    if light_rows:
        for n, row in enumerate(table.rows[1:]):
            if n % 2 == 1:
                for c in row.cells:
                    shade(c, LIGHT_HEX)
    for row in table.rows[1:]:
        for c in row.cells:
            for p in c.paragraphs:
                for r in p.runs:
                    r.font.size = Pt(9)


def add_confidential_footer(doc) -> None:
    """Mismo pie que cyber_health_report.py: 'Confidencial · Pagina N'
    centrado, en gris, con el campo PAGE nativo de Word (se actualiza
    solo al paginar, no es un numero fijo escrito a mano)."""
    sec = doc.sections[0]
    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
    fr = fp.add_run("Confidencial · Página ")
    fr.font.size = Pt(8)
    fr.font.color.rgb = MUTED
    run = fp.add_run()
    for tag, txt in (("begin", None), (None, "PAGE"), ("end", None)):
        if tag:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), tag)
        else:
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = txt
        run._r.append(el)


